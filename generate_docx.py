import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_proposal():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(12)

    # Title
    title = doc.add_heading('具身智能实验室构建项目书', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. 项目概述
    doc.add_heading('一、 项目概述', level=1)
    doc.add_heading('1.1 建设背景与理念', level=2)
    doc.add_paragraph('本项目的核心目标是为具身智能研究建立一套可持续运转的“实验室运行框架（Lab Operating Framework）”。有别于传统的单一演示场景开发，本项目致力于打造标准化的物理空间、统一的软硬件接口、严密的安全规范以及清晰的组织协作机制。通过“框架优先、共用复用”的建设理念，为后续各类具身智能前沿课题（如跨形态协同、遥操作蒸馏等）提供坚实、可扩展的系统底座。')

    doc.add_heading('1.2 建设目标与边界', level=2)
    doc.add_paragraph('本项目旨在实现实验室的“可运营、可纳管、可复用、可审计与可扩展”。在框架建设期内，项目边界严格限定于基础设施落地、已有5台异构机器人（双足、四足、机械臂等）的平台纳管、共用中间件部署及管理制度的建立，暂不涉及具体的业务场景联调与算法模型训练。')

    doc.add_heading('1.3 核心运行约束', level=2)
    doc.add_paragraph('为确保工程落地与资源可控，项目设定以下核心运行基线：')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('并发与规模约束：').bold = True
    p.add_run(' 框架期最大支持2台机器人同时进行动态实验，且所有动态实验必须在指定的单一物理安全分区内进行。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('资源与投资约束：').bold = True
    p.add_run(' 实验室构建开发总投资额度严格控制在20万元人民币以内，统筹用于网络、服务器及安全设施建设。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('数据与安全底线：').bold = True
    p.add_run(' 动态实验必须全量录制数据并绑定唯一实验批次号（Run ID）；现场必须保证至少2人在场，且具备直接切断动力电源的物理急停能力。')

    # 2. 系统总体架构设计
    doc.add_heading('二、 系统总体架构设计', level=1)
    doc.add_paragraph('系统架构采用自下而上的六层模型设计，确保软硬解耦与模块化协作。')

    doc.add_heading('2.1 六层架构模型', level=2)
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '层级'
    hdr_cells[1].text = '核心内容与职责'
    data = [
        ('L6 组织与流程', '职责矩阵 (RACI)、安全操作规程 (SOP)、实验与变更流程'),
        ('L5 系统流水线', '感知 (Perception) → 规划 (Plan) → 技能 (Skill) → 安全 (Safety)'),
        ('L4 数据与实验', '实验批次号 (Run ID)、数据录制、索引与元数据管理'),
        ('L3 软件中间件', 'ROS2 Humble、DDS 通信、统一消息接口'),
        ('L2 硬件与运维', '资产台账、设备点检、驱动桥 (Driver Bridge)'),
        ('L1 空间与设施', '场地分区、网络拓扑、算力服务器、物理安全设施')
    ]
    for i, (layer, desc) in enumerate(data):
        row_cells = table.rows[i+1].cells
        row_cells[0].text = layer
        row_cells[1].text = desc

    doc.add_heading('2.2 统一流水线与软硬解耦', level=2)
    doc.add_paragraph('为解决异构硬件适配问题，所有底层硬件SDK均被封装于“驱动桥”之下，对上层暴露统一的通用技能指令。系统流水线严格切分感知、规划、技能与安全模块的输入输出语义。特别地，安全（Safety）模块作为独立节点置于流水线最末端，具备“一票否决权”，一旦触发红线可直接覆盖上层指令执行软停。')

    # 3. 基础设施与物理空间规划
    doc.add_heading('三、 基础设施与物理空间规划', level=1)
    doc.add_heading('3.1 场地平面布局与安全分区', level=2)
    doc.add_paragraph('实验室物理空间划分为四大核心功能区，实现人机物理隔离：')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('动态实验区 (Z-DYN)：').bold = True
    p.add_run(' 采用2米高工业铝型材安全围栏封闭，铺设防静电地胶，专供机器人全速动态测试。实验期间严禁人员入内。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('操作与监控区 (Z-OP)：').bold = True
    p.add_run(' 位于围栏外部，部署主控工作站，配备主物理急停控制箱，为绝对安全区。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('维护与充电区 (Z-MAINT)：').bold = True
    p.add_run(' 用于机器人日常维修、标定及电池充电，配备专用消防器材。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('设备存放区 (Z-STORE)：').bold = True
    p.add_run(' 用于存放处于断电闲置状态的机器人。')

    doc.add_heading('3.2 算力与网络拓扑规划', level=2)
    doc.add_paragraph('为满足端到端控制延迟小于10ms的非功能性需求，实验室构建独立的万兆核心局域网与专用Wi-Fi 6无线网络，与日常办公网络物理隔离。算力分配方面，主节点工作站负责轻量级规划与高实时性控制，重负载工作站负责视觉大模型推理、离线仿真与数据录制。')

    doc.add_heading('3.3 软硬安全接口与电气逻辑', level=2)
    doc.add_paragraph('鉴于机器人均采用电池供电，安全系统采用“全局软件急停+遥控器硬件阻断+本体物理断电”三级保障机制。操作台与安全门入口处的急停按钮接入主控站GPIO，触发全网软件级急停广播；动态实验时安全员必须手持机器人专属遥控器，遇紧急情况人为触发遥控器硬件掉电开关；在确保人员安全的前提下，可直接按下机器人本体的物理急停按钮切断电池供电。')

    # 4. 软件平台与数据规范
    doc.add_heading('四、 软件平台与数据规范', level=1)
    doc.add_heading('4.1 统一消息接口与版本基线', level=2)
    doc.add_paragraph('系统底层基于 ROS2 Humble 构建，采用 Cyclone DDS 保障通信质量。所有计算节点操作系统统一锁定为 Ubuntu 22.04 LTS。消息接口方面，严格定义了感知、规划、技能与安全四大类 Topic 的命名空间、消息类型及服务质量 (QoS) 策略，确保跨设备通信的标准化。')

    doc.add_heading('4.2 数据落盘与存储策略', level=2)
    doc.add_paragraph('实验数据采用冷热分层存储架构。基于2台设备并发、日均产生约1.5TB数据的基线评估，近期热数据（0-7天）存储于重负载工作站的NVMe固态硬盘中；历史冷数据则通过自动化脚本定期迁移至大容量NAS存储阵列。每次实验强制生成包含时间、操作人、设备及代码版本哈希值的唯一 Run ID，确保实验数据的100%血缘追溯。')

    # 5. 实验室运营与管理体系
    doc.add_heading('五、 实验室运营与管理体系', level=1)
    doc.add_heading('5.1 组织架构与职责划分', level=2)
    doc.add_paragraph('项目组由三人核心团队构成，职责边界清晰：')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('系统/规划负责人 (P1)：').bold = True
    p.add_run(' 统筹总体方案、系统架构、数据流契约及总预算审批。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('平台/基础设施工程师 (P2)：').bold = True
    p.add_run(' 负责网络算力落地、ROS2接口实现、代码仓库及CI/CD流水线建设。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('现场/资产/安全工程师 (P3)：').bold = True
    p.add_run(' 负责场地分区、物理安全设施、资产台账及安全操作规程的落地执行。')

    doc.add_heading('5.2 实验标准操作规程 (SOP)', level=2)
    doc.add_paragraph('制定了严格的实验五阶段流程（申请、准备、执行、复盘、归档）。明确了“二人规则”红线，即中高风险实验必须两人在场，一人操作，一人值守物理急停。规定了机器人失控及电池起火等紧急情况的标准处置预案。')

    doc.add_heading('5.3 变更控制与设备维护', level=2)
    doc.add_paragraph('设立每周五的固定变更评审与复盘会议。任何涉及架构、接口、安全策略及预算的变更均需提交变更请求（CR）并经相应负责人审批。同时，建立了每日5分钟例行点检与周度深度维护制度，赋予点检人员发现隐患时直接冻结设备的权限。')

    # 6. 项目实施计划与预算
    doc.add_heading('六、 项目实施计划与预算', level=1)
    doc.add_heading('6.1 建设分期与里程碑', level=2)
    doc.add_paragraph('项目整体规划为16周，分为四个核心阶段：')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('阶段A：框架定义 (W1-W4)。').bold = True
    p.add_run(' 完成总体方案、架构蓝图、详细设计与采购下单。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('阶段B：基础设施落地 (W5-W10)。').bold = True
    p.add_run(' 完成网络布线、围栏安装、服务器上架及基础设施联合验收。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('阶段C：平台上线 (W11-W14)。').bold = True
    p.add_run(' 完成中间件部署、数据流打通、CI流水线建设及软硬安全联合演练。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('阶段D：框架验收 (W15-W16)。').bold = True
    p.add_run(' 执行标准实验流程交叉验证，达成框架总验收，正式对外提供项目接入能力。')

    doc.add_heading('6.2 投资预算明细', level=2)
    doc.add_paragraph('首期建设预算申请额度约为 4.85 万元人民币，远低于 20 万元的总投资上限。资金主要用于：')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('网络与存储：').bold = True
    p.add_run(' 万兆核心交换机、企业级Wi-Fi 6 AP、4盘位NAS及企业级硬盘（约 1.85 万元）。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('场地与安全：').bold = True
    p.add_run(' 工业铝型材安全围栏、物理急停控制箱、防静电地胶及消防器材（约 2.00 万元）。')
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('备件与易耗品：').bold = True
    p.add_run(' 机器人备用电池、高精度标定板及日常耗材（约 1.00 万元）。')

    output_path = os.path.join(os.getcwd(), '具身智能实验室构建项目书.docx')
    doc.save(output_path)
    print(f"Successfully saved to {output_path}")

if __name__ == '__main__':
    create_proposal()
